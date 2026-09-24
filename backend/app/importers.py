"""Leitura e validação das planilhas recebidas pela área administrativa."""

from datetime import date, datetime
from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Promoter, Store


def _header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _identifier(cell) -> str:
    """Preserva zeros à esquerda quando a formatação do Excel os define."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = int(value) if float(value).is_integer() else value
        zeros = re.fullmatch(r"0+", cell.number_format or "")
        if zeros and isinstance(number, int):
            return f"{number:0{len(zeros.group())}d}"
        return str(number)
    return _text(value)


def _date(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), pattern)
            except ValueError:
                pass
    return None


def _join(*values: object) -> str:
    return ", ".join(value for value in (_text(item) for item in values) if value)


""" def _format_address(
    street: object,
    number_or_complement: object,
    neighborhood: object,
    city: object,
    state: object,
    cep: object = None,
) -> str:
    street_part = _join(street, number_or_complement)

    city_text = _text(city)
    state_text = _text(state).upper()

    city_state = (
        f"{city_text}/{state_text}"
        if city_text and state_text
        else city_text or state_text
    )

    location_part = (
        f"{_text(neighborhood)} - {city_state}"
        if _text(neighborhood) and city_state
        else _text(neighborhood) or city_state
    )

    address = " - ".join(
        part for part in [street_part, location_part] if part
    )

    cep_text = _text(cep)

    if cep_text:
        address = f"{address} - CEP: {cep_text}"

    return address """


def _rows(file_bytes: bytes, required_headers: set[str]):
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    first_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
    if not first_row:
        raise ValueError("A planilha está vazia.")
    columns = {_header(cell.value): index for index, cell in enumerate(first_row)}
    missing = sorted(required_headers - columns.keys())
    if missing:
        raise ValueError(f"Colunas obrigatórias ausentes: {', '.join(missing)}.")
    return sheet.iter_rows(min_row=2), columns


def import_promoters(db: Session, file_bytes: bytes) -> dict:
    rows, columns = _rows(
        file_bytes,
        {
            "MATRICULA",
            "FUNCIONARIO",
            "DATA DE ADMISSAO",
            "CARGO",
            "SITUACAO",
            "ENDERECO",
            "BAIRRO",
            "MUNICIPIO",
            "ESTADO",
            "CEP",
        },
    )

    prepared, errors = [], []

    for line, row in enumerate(rows, start=2):
        value = lambda header: row[columns[header]]

        employee_id = _identifier(value("MATRICULA"))
        name = _text(value("FUNCIONARIO").value)
        cpf = _identifier(value("CPF")) if "CPF" in columns else ""
        admission_date = _date(value("DATA DE ADMISSAO").value)

        if not any(cell.value is not None for cell in row):
            continue

        problems = []

        if not employee_id:
            problems.append("Matrícula não informada")

        if not name:
            problems.append("Funcionário não informado")

        if not cpf:
            problems.append("CPF não informado")

        if not admission_date:
            problems.append("Data de Admissão inválida ou não informada")

        if problems:
            errors.append({
                "line": line,
                "message": "; ".join(problems),
            })
            continue

        get = lambda header: value(header).value if header in columns else None

        prepared.append({
            "employee_id": employee_id,
            "name": name,
            "cpf": cpf,
            "rg": _identifier(value("RG")) if "RG" in columns else "",
            "ctps": (
            f"{_identifier(value('NO CARTEIRA PROF.'))}/ "
            f"{_identifier(value('SERIE CP'))}"
            if "NO CARTEIRA PROF." in columns and "SERIE CP" in columns
            else ""
            ),
            "pis": _identifier(value("PIS")) if "PIS" in columns else "",
            "address": _text(get("ENDERECO")) or None,
            "neighborhood": _text(get("BAIRRO")) or None,
            "city": _text(get("MUNICIPIO")) or None,
            "state": _text(get("ESTADO")).upper()[:2] or None,
            "cep": (
                _identifier(value("CEP"))
                if "CEP" in columns
                else None
            ),
            "admission_date": admission_date,
            "position": _text(get("CARGO")) or "PROMOTER",
            "active": _header(get("SITUACAO")) in {"ATIVO", "ATIVA"},
        })

    if errors:
        return {
            "created": 0,
            "updated": 0,
            "errors": errors,
        }

    created = 0
    updated = 0

    try:
        for data in prepared:
            promoter = (
                db.query(Promoter)
                .filter(Promoter.cpf == data["cpf"])
                .first()
            )

            if promoter:
                for field, field_value in data.items():
                    setattr(promoter, field, field_value)

                updated += 1
            else:
                db.add(Promoter(**data))
                created += 1

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "created": created,
        "updated": updated,
        "errors": [],
    }


def import_stores(db: Session, file_bytes: bytes) -> dict:
    rows, columns = _rows(
        file_bytes,
        {
            "COD. CLIENTE",
            "NOME REDUZIDO",
            "END. CLIENTE",
            "BAIRRO",
            "CIDADE",
            "UF",
            "CEP",
        },
    )

    prepared, errors = [], []

    for line, row in enumerate(rows, start=2):
        value = lambda header: row[columns[header]]

        code = _identifier(value("COD. CLIENTE"))
        name = _text(value("NOME REDUZIDO").value)

        if not any(cell.value is not None for cell in row):
            continue

        if not code or not name:
            missing = (
                "Código da loja"
                if not code
                else "Nome reduzido"
            )

            errors.append({
                "line": line,
                "message": f"{missing} não informado",
            })

            continue

        get = lambda header: (
            value(header).value
            if header in columns
            else None
        )

        prepared.append({
            "code": code,
            "name": name,
            "client": _text(get("NOME CLIENTE")) or None,

            # Endereço com rua e número.
            "address": _text(get("END. CLIENTE")) or None,

            # Bairro separado.
            "neighborhood": _text(get("BAIRRO")) or None,

            # Cidade e UF separados.
            "city": _text(get("CIDADE")) or None,
            "state": _text(get("UF")).upper()[:2] or None,

            "active": True,
        })
        
    if errors:
        return {
            "created": 0,
            "updated": 0,
            "errors": errors,
        }

    created = 0
    updated = 0

    try:
        for data in prepared:
            store = (
                db.query(Store)
                .filter(Store.code == data["code"])
                .first()
            )

            if store:
                for field, field_value in data.items():
                    setattr(store, field, field_value)

                updated += 1

            else:
                db.add(Store(**data))
                created += 1

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "created": created,
        "updated": updated,
        "errors": [],
    }
